from pathlib import Path
from math import pi, atan, tan, degrees
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

src = Path(r"<downloads>/Plantilla_Calculo_Engranajes.xlsx")
out = Path(__file__).parent / "Plantilla_Calculo_Engranajes_STUDENT_C.xlsx"
wb = load_workbook(src)

for sheet, n in [(wb["Menos_30"], 19), (wb["30_o_mas"], 38)]:
    sheet["B4"] = n
    sheet["B5"] = 20
    sheet["B6"] = 2

if "RESUMEN_STUDENT_C" in wb.sheetnames:
    del wb["RESUMEN_STUDENT_C"]
ws = wb.create_sheet("RESUMEN_STUDENT_C", 0)
ws.append(["STUDENT C", "Z1 ≥30", "Z2 <30"])
ws.append(["Variable", "N=38", "N=19"])
values = [
    ("N",38,19),("M",2,2),("Alfa",20,20),("Ángulo hélice",25,-25),
    ("De",80,42),("Dp",76,38),("Ct",71.42,35.71),("Di",71,33.33),
    ("Par/Pitch",6.28,6.28),("Delta",2.37,4.74),("R1",9.5,2.58),
    ("R2","No aplica",8.13),("R3","No aplica",12.2),("r",0.6,0.6),
    ("Acara",8,8),("Chaflán",1.56,0.8),("P/4",1.57,"No aplica"),
    ("Distancia entre ejes",57,57),("Giro posterior θ",5.33,10.07),("Giro ±θ/2",2.66,5.04),
]
for row in values: ws.append(row)
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF", size=12)
    cell.fill = PatternFill("solid", fgColor="17324D")
for cell in ws[2]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="246B9E")
for col, width in {"A":28,"B":18,"C":18}.items(): ws.column_dimensions[col].width=width
ws.freeze_panes="A3"
ws.sheet_view.showGridLines=False
for row in ws.iter_rows():
    for cell in row: cell.alignment=Alignment(vertical="center")

try:
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
except Exception:
    pass
wb.save(out)
print(out)
